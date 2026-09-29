import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_proempresa_template():
    doc = Document()
    
    # Configure styles
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(10)
    
    # Date
    p = doc.add_paragraph('Lima, [FECHA_ACTUAL]')
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    # Title
    p = doc.add_paragraph('¡ NOTIFICACIÓN DE COBRANZA JUDICIAL!')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(12)
    
    # Header Data
    doc.add_paragraph('SEÑOR (A) (ES): [NOMBRE_CLIENTE]')
    doc.add_paragraph('CUENTA: [NUM_CUENTA]')
    doc.add_paragraph('CONYUGE: [CONYUGE_CLIENTE]')
    doc.add_paragraph('DOMICILIO: [DIRECCION_CLIENTE]')
    doc.add_paragraph('AVAL: [NOMBRE_AVAL]')
    doc.add_paragraph('DOMICILIO AVAL: [DOMICILIO_AVAL]')
    doc.add_paragraph('AGENCIA: [AGENCIA]')
    
    # Body
    doc.add_paragraph()
    doc.add_paragraph('Estimado Cliente:')
    p = doc.add_paragraph('Para saludarlo(s) a nombre del Departamento de Recuperaciones de ProEmpresa y comunicarles que se le brinda la oportunidad de solucionar su crédito vía negociación dentro de las políticas y facilidades que otorga la Financiera, para cuyo efecto lo invitamos a acercarse dentro de las 48 horas de recibo la presente comunicación a la Agencia de ProEmpresa donde obtuvo su crédito y contactarse con el Administrador.')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    # Debt Box
    doc.add_paragraph('CANCELACIÓN HASTA EL:                      ')
    table = doc.add_table(rows=2, cols=1)
    table.style = 'Table Grid'
    cell0 = table.cell(0, 0)
    cell0.text = 'MONTO TOTAL DE LA DEUDA'
    cell0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell0.paragraphs[0].runs[0].bold = True
    
    cell1 = table.cell(1, 0)
    cell1.text = 'S/ [MONTO_DEUDA]'
    cell1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell1.paragraphs[0].runs[0].bold = True
    
    # Warning
    doc.add_paragraph()
    p = doc.add_paragraph('Se le recuerda que “todo pago se realiza en las agencias de Financiera ProEmpresa o en las cuentas que mantiene está en otros bancos, nunca a nombre de personas particulares y menos entrega de efectivo a terceros”.')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    # Footer
    doc.add_paragraph()
    doc.add_paragraph('Atentamente,')
    doc.add_paragraph()
    doc.add_paragraph('Dra. Lorena Prada Gonzales                                       986225114')
    doc.add_paragraph('Abogada ProEmpresa')
    doc.add_paragraph('---------------------------------------------------------')
    
    # Meta
    doc.add_paragraph('Negocio: [DIR_NEGOCIO]     Tipo de PROCESO: [TIPO_PROCESO]')
    doc.add_paragraph('Teléfono de cliente: [TELEFONO_CLIENTE]')
    doc.add_paragraph('Fecha de pago: [FECHA_ULT_PAGO]')
    
    # Save
    doc.save(r'c:\CRM PYP\plantilla_proempresa.docx')
    print("TEMPLATE CREADO EXITOSAMENTE")

if __name__ == '__main__':
    create_proempresa_template()
