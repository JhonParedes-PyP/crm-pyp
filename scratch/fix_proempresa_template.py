from docx import Document

def fix_template():
    doc = Document(r'c:\CRM PYP\MODELO DE CARTA DE PROEMPRESA.docx')
    
    replacements = {
        "MERINO TRILLO ILBER": "[NOMBRE_CLIENTE]",
        "098785010045895941": "[NUM_CUENTA]",
        "CARRASCO RUIZ JACQUELINE ASTRID": "[CONYUGE_CLIENTE]",
        "CALLE LOS FRUTALES Lt 55 ASENTAMIENTO HUMANO HUAYCAN ETAPA 4 ZONA N UCV 175 - ATE VITARTE - LIMA - LIMA ATE": "[DIRECCION_CLIENTE]",
        "AGENCIA COMERCIAL": "[AGENCIA]",
        "27/01/2026": "[FECHA_ULT_PAGO]",
        "Lima          , de                 de 2026": "Lima, [FECHA_ACTUAL]"
    }
    
    for p in doc.paragraphs:
        # Also fix the AVAL lines manually since they are empty in their doc
        if p.text.strip() == "AVAL:":
            p.text = "AVAL: [NOMBRE_AVAL]"
        if p.text.strip() == "DOMICILIO:":
            p.text = "DOMICILIO AVAL: [DOMICILIO_AVAL]"
            
        for k, v in replacements.items():
            if k in p.text:
                p.text = p.text.replace(k, v)
                
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for k, v in replacements.items():
                        if k in p.text:
                            p.text = p.text.replace(k, v)
                            
    # also we need to add the bottom data that Proempresa requested
    # Negocio: [DIR_NEGOCIO]     Tipo de PROCESO: [TIPO_PROCESO]
    # Teléfono de cliente: [TELEFONO_CLIENTE]
    # Check if they exist
    has_meta = any("[DIR_NEGOCIO]" in p.text for p in doc.paragraphs)
    if not has_meta:
        doc.add_paragraph('Negocio: [DIR_NEGOCIO]     Tipo de PROCESO: [TIPO_PROCESO]')
        doc.add_paragraph('Teléfono de cliente: [TELEFONO_CLIENTE]')
        doc.add_paragraph('Fecha de pago: [FECHA_ULT_PAGO]')
        
    doc.save(r'c:\CRM PYP\MODELO DE CARTA DE PROEMPRESA.docx')
    print("TEMPLATE FIXED")

if __name__ == "__main__":
    fix_template()
