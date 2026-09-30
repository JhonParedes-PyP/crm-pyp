from docx import Document
from docx.shared import Pt

def format_text(p, text, size, bold=True):
    p.clear()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)

def fix_proempresa_template():
    doc = Document(r'c:\CRM PYP\MODELO DE CARTA DE PROEMPRESA.docx')
    
    # 1. Change Lima, [FECHA_ACTUAL] to Lima,
    for p in doc.paragraphs:
        if "Lima, [FECHA_ACTUAL]" in p.text:
            p.text = p.text.replace("Lima, [FECHA_ACTUAL]", "Lima,                                           ")
            # preserve alignment if needed, typically right-aligned. We just replaced text.
            
    # 2. Update font sizes and bold
    # The client is at paragraph 4 and 34
    # The spouse is at paragraph 6 and 36
    # Instead of hardcoding indices, we can search for the start strings.
    for p in doc.paragraphs:
        if p.text.startswith("SEÑOR (A) (ES):") or p.text.startswith("SEOR (A) (ES):"):
            format_text(p, p.text, 14, True)
        elif p.text.startswith("CONYUGE:"):
            format_text(p, p.text, 12, True)
            
    # 3. Delete the hardcoded cargo info from paragraphs 55, 56, 57
    # We will identify them by content and remove them.
    # Actually, we can just clear their text to remove them.
    for p in doc.paragraphs:
        if "Negocio: CALLE LOS FRUTALES" in p.text:
            p.text = ""
        if "Teléfono de cliente: 991823967" in p.text or "Telfono de cliente: 991823967" in p.text:
            p.text = ""
        if "Fecha de pago: 7/13/2026" in p.text:
            p.text = ""
            
    # 4. We should also remove some empty lines before the final dynamic info so it fits on the page.
    # We'll just remove a few empty paragraphs near the end.
    empty_removed = 0
    for p in reversed(doc.paragraphs):
        if not p.text.strip() and empty_removed < 4:
            p._element.getparent().remove(p._element)
            empty_removed += 1
            
    doc.save(r'c:\CRM PYP\MODELO DE CARTA DE PROEMPRESA.docx')
    print("TEMPLATE UPDATED SUCCESSFULLY")

if __name__ == "__main__":
    fix_proempresa_template()
